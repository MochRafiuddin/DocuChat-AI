import gradio as gr
from PyPDF2 import PdfReader
import docx
from pptx import Presentation
from openai import OpenAI
from dotenv import load_dotenv
import os
import time  # Tambahkan di atas jika belum ada
load_dotenv()

api_key = os.environ["OPENAI_API_KEY"]
client = OpenAI(api_key=api_key)

def read_file(file):
    if file.name.endswith(".pdf"):
        reader = PdfReader(file.name)
        return "\n".join(page.extract_text() for page in reader.pages if page.extract_text())
    elif file.name.endswith(".docx"):
        doc = docx.Document(file.name)
        return "\n".join(p.text for p in doc.paragraphs)
    elif file.name.endswith(".pptx"):
        prs = Presentation(file.name)
        return "\n".join(shape.text for slide in prs.slides for shape in slide.shapes if hasattr(shape, "text"))
    elif file.name.endswith(".txt"):
        return file.read().decode("utf-8")
    return ""

all_text = ""

def upload_and_store(files):
    global all_text
    all_text = ""
    for file in files:
        all_text += read_file(file) + "\n\n"
    return "✅ Uploaded and combined all content."

def ask_question(question):
    start_embed = time.time()
    global all_text
    if not all_text:
        return "❌ Please upload files first."
    
    prompt = f"You are an assistant. Based on the following content:\n\n{all_text[:100000]}\n\nAnswer this question:\n{question}"
    
    response = client.chat.completions.create(
        # model="gpt-4o-mini-2024-07-18",
        model="gpt-4.1-nano-2025-04-14",
        messages=[{"role": "user", "content": prompt}],
        temperature=0,
    )
    print(f"✅ pertanyaan selesai dalam {time.time() - start_embed:.2f} detik")
    return response.choices[0].message.content.strip()

with gr.Blocks() as demo:
    gr.Markdown("### 🧠 Chat with Uploaded Files (No Embedding)")
    files = gr.File(file_types=[".pdf", ".docx", ".pptx", ".txt"], file_count="multiple")
    upload_btn = gr.Button("Upload Files")
    status = gr.Textbox(label="Status")
    q = gr.Textbox(label="Your Question")
    a = gr.Textbox(label="Answer", lines=5)
    ask_btn = gr.Button("Ask")

    upload_btn.click(upload_and_store, inputs=files, outputs=status)
    ask_btn.click(ask_question, inputs=q, outputs=a)

demo.launch()